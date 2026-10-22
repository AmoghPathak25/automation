import pandas as pd
from openpyxl import load_workbook
from src.readers.excel_reader import read_tabular
from src.reporting.excel_exporter import export_excel
from src.services.normalizer import normalize_frame
from src.services.validator import validate_inventory

def test_synthetic_workbook_validation_and_roundtrip(tmp_path):
    source = tmp_path / "synthetic_sap_like.xlsx"
    rows = [
        {"Material Number": "000041", "Material Description": "Synthetic item", "Unrestricted Stock": 12, "Stock Date": "2025-01-05"},
        {"Material Number": "000042", "Material Description": "Missing quantity", "Unrestricted Stock": None, "Stock Date": "2025-01-06"},
        {"Material Number": "000043", "Material Description": "Negative quantity", "Unrestricted Stock": -2, "Stock Date": "not-a-date"},
        {"Material Number": "000043", "Material Description": "Duplicate", "Unrestricted Stock": 4, "Stock Date": "2025-01-07"},
    ]
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="Stock Export", index=False)
    before = source.read_bytes()
    frame, sheet = read_tabular(source, "Stock Export")
    normalized = normalize_frame(frame, source.name, sheet)
    result = validate_inventory(normalized, date_format="%Y-%m-%d")
    assert result["material_code"].tolist() == ["000041", "000042", "000043", "000043"]
    assert result.iloc[1]["quantity_on_hand"] is None
    assert result.iloc[1]["validation_status"] == "NEEDS_REVIEW"
    assert any("Negative" in issue for issue in result.iloc[2]["validation_errors"])
    assert any("Invalid stock_date" in issue for issue in result.iloc[2]["validation_errors"])
    assert any("Duplicate" in issue for issue in result.iloc[2]["validation_errors"])
    assert any("Duplicate" in issue for issue in result.iloc[3]["validation_errors"])
    assert source.read_bytes() == before
    output = export_excel(result, tmp_path / "validation-output.xlsx", "Validated Inventory")
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Validated Inventory"]
    sheet_out = workbook["Validated Inventory"]
    headers = [cell.value for cell in sheet_out[1]]
    assert "material_code" in headers and "validation_status" in headers and "source_row" in headers
    code_col = headers.index("material_code") + 1
    assert [sheet_out.cell(row=i, column=code_col).value for i in range(2, 6)] == ["000041", "000042", "000043", "000043"]
    assert source.read_bytes() == before

def test_repeat_import_export_has_stable_row_count_and_no_source_mutation(tmp_path):
    source = tmp_path / "repeat.xlsx"
    pd.DataFrame({"Material Number": ["0009", "0010"], "Unrestricted Stock": [1, 2]}).to_excel(source, index=False)
    original_bytes = source.read_bytes()
    exports = []
    for n in range(2):
        frame, sheet = read_tabular(source)
        records = validate_inventory(normalize_frame(frame, source.name, sheet))
        out = export_excel(records, tmp_path / "output.xlsx")
        exports.append(pd.read_excel(out, dtype={"material_code": str}))
    assert len(exports[0]) == len(exports[1]) == 2
    assert exports[0]["material_code"].tolist() == exports[1]["material_code"].tolist() == ["0009", "0010"]
    assert source.read_bytes() == original_bytes
