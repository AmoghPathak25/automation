from pathlib import Path
import os
import sys
from streamlit.testing.v1 import AppTest
import pandas as pd
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
SECTIONS = ["Overview", "Import files", "Validation report", "Inventory search", "Clearance candidates", "Email request review", "Offer draft review", "Processing history", "Configuration and rules"]

def test_every_dashboard_section_renders_without_exception():
    for section in SECTIONS:
        app = AppTest.from_file(str(ROOT / "src" / "app.py"), default_timeout=15).run()
        app.radio[0].set_value(section).run()
        assert not app.exception, f"{section} raised Streamlit exceptions: {app.exception}"
        assert app.title and app.title[0].value == "Office Email Automation"


def test_widget_text_saves_immediately_and_survives_app_rerun(tmp_path, monkeypatch):
    import src.database.repository as repository_module
    repository = repository_module.Repository(tmp_path / "ui_drafts.sqlite3")
    monkeypatch.setattr(repository_module, "Repository", lambda: repository)
    app = AppTest.from_file(str(ROOT / "src" / "app.py"), default_timeout=15).run()
    app.session_state["email_request"] = {
        "request_key": "synthetic-email-key", "sender": "customer@example.invalid",
        "subject": "Synthetic request", "category": "Product availability enquiry",
        "material_code": "SYN-0001", "grade": None, "dimensions": None,
        "requested_quantity": None, "uncertainties": ["synthetic fixture"], "confidence": "LOW"
    }
    app.radio[0].set_value("Email request review").run()
    draft_fields = [field for field in app.text_area if field.label.startswith("Unsent draft")]
    assert len(draft_fields) == 1
    exact_text = "Edited and saved on the same Streamlit interaction."
    draft_fields[0].set_value(exact_text)
    save_button = next(button for button in app.button if button.label == "Save draft for human review")
    save_button.click().run()
    saved = repository.drafts()
    assert len(saved) == 1
    assert saved[0]["payload"]["body"] == exact_text
    app.run()
    persisted_field = next(field for field in app.text_area if field.label.startswith("Unsent draft"))
    assert persisted_field.value == exact_text

def test_upload_import_persist_export_streamlit_workflow(tmp_path, monkeypatch):
    import sqlite3
    import src.database.repository as repository_module
    from src.reporting.excel_exporter import export_excel
    from src.services.normalizer import normalize_frame
    from src.services.validator import validate_inventory
    from src.utils.file_safety import sha256_file
    from src.utils.upload_staging import StagedUpload
    source = tmp_path / "synthetic_workbook.xlsx"
    pd.DataFrame({"Material Number": ["000321", "000322"], "Material Description": ["Synthetic item A", "Synthetic item B"], "Unrestricted Stock": [7, None], "Stock Date": ["2026-01-15", "ambiguous"]}).to_excel(source, index=False)
    source_copy = source.read_bytes()
    database = tmp_path / "workflow.sqlite3"
    repository = repository_module.Repository(database)
    monkeypatch.setattr(repository_module, "Repository", lambda: repository)
    staged_root = tmp_path / "staging"
    with StagedUpload(staged_root, source.name, source.read_bytes()) as staged:
        frame, sheet = __import__("src.readers.excel_reader", fromlist=["read_tabular"]).read_tabular(staged.path)
        normalized = normalize_frame(frame, source.name, sheet)
        records = validate_inventory(normalized, date_format="%Y-%m-%d")
        run_id = repository.save_import(source.name, sha256_file(staged.path), sheet, records, records.attrs["summary"])
        export_path = export_excel(records, tmp_path / "workflow_report.xlsx", "Validation", run_id=run_id)
        assert export_path.exists()
    assert source.read_bytes() == source_copy
    persisted = {record["material_code"]: record for record in repository.records(run_id)}
    assert len(persisted) == 2
    assert persisted["000321"]["material_code"] == "000321"
    assert persisted["000322"]["validation_status"] == "NEEDS_REVIEW"
    with sqlite3.connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 2
    from openpyxl import load_workbook
    workbook = load_workbook(export_path, read_only=True, data_only=True)
    assert workbook.active.max_row == 3
    assert "000321" in [cell[0] for cell in workbook.active.iter_rows(min_row=2, values_only=True)]
    app = AppTest.from_file(str(ROOT / "src" / "app.py"), default_timeout=15).run()
    assert not app.exception
    assert app.title[0].value == "Office Email Automation"
    app.radio[0].set_value("Validation report").run()
    assert not app.exception
